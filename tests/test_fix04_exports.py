"""FIX-04: complete CSV/XLSX exports, XLSX policy, spreadsheet safety, memory."""
import csv
import glob
import io
import json
import os
import tempfile
import tracemalloc

import duckdb
import pytest
from openpyxl import load_workbook

from fix01_corpus import make_client

CATS = ["Total", "a,b", "line\nbreak", "=SUM(1)", 'Țară "x"', "-5", "+cmd", "@x"]
CAT_EN = {"Total": "Total EN", "Țară \"x\"": "Country EN"}
AREAS = ["Cluj", "Iasi", "Bihor", "Arad", "Ilfov"]
N_PERIODS = 1600          # 8 * 5 * 1600 = 64,000 rows (> 50,000 chart cap)
N_ROWS = len(CATS) * len(AREAS) * N_PERIODS


def _periods():
    return [f"{2000 + i // 12}-{i % 12 + 1:02d}" for i in range(N_PERIODS)]


_ROWS_CACHE = []


def _source_rows():
    if _ROWS_CACHE:
        return _ROWS_CACHE
    rows = _ROWS_CACHE
    k = 0
    for c in CATS:
        for a in AREAS:
            for t in _periods():
                k += 1
                v = None if k % 997 == 0 else (k % 100 if k % 3 == 0 else k / 7)
                rows.append((c, a, t, v))
    return rows


@pytest.fixture
def client(tmp_path, monkeypatch):
    c = make_client(tmp_path, monkeypatch)
    corpus = tmp_path / "corpus"
    import pyarrow as pa
    import pyarrow.parquet as pq
    cols = list(zip(*_source_rows()))
    pq.write_table(pa.table({"CATEGORY": cols[0], "REF_AREA": cols[1],
                             "TIME_PERIOD": cols[2],
                             "OBS_VALUE": pa.array(cols[3], pa.float64())}),
                   corpus / "parquet" / "BIG1.parquet")
    db = duckdb.connect(str(corpus / "metadata.duckdb"))
    db.execute("CREATE TABLE sdmx_codes (nom_item_id INTEGER, display_label_en VARCHAR)")
    db.execute("INSERT INTO matrices (matrix_code, matrix_name, matrix_name_en, row_count) "
               "VALUES ('BIG1', 'Big', 'Big', ?)", [N_ROWS])
    for i, (label, col) in enumerate(
            [("Category", "CATEGORY"), ("Area", "REF_AREA"), ("Time", "TIME_PERIOD")], 1):
        db.execute("INSERT INTO dimensions VALUES (?, 'BIG1', ?, ?, ?, 3)",
                   [1000 + i, i, label, col])
    for n, c_ in enumerate(CATS, 1):
        db.execute("INSERT INTO dimension_options VALUES (1001, ?, ?)", [100 + n, c_])
        if c_ in CAT_EN:
            db.execute("INSERT INTO sdmx_codes VALUES (?, ?)", [100 + n, CAT_EN[c_]])
    db.close()
    return c


def dl(client, code="BIG1", **p):
    return client.get(f"/api/datasets/{code}/download", params=p)


def parse_csv(resp):
    return list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))


def key(r):
    return (r[0], r[1], r[2])


def sel(rows, cats=None, periods=None):
    out = [r for r in rows
           if (cats is None or r[0] in cats) and (periods is None or r[2] in periods)]
    return out


# ------------------------------------------------------------ 1. completeness

def test_csv_exports_every_row_beyond_chart_cap(client):
    r = dl(client)
    assert r.status_code == 200
    rows = parse_csv(r)
    assert rows[0] == ["CATEGORY", "REF_AREA", "TIME_PERIOD", "OBS_VALUE"]
    body = rows[1:]
    assert len(body) == N_ROWS > 50_000
    assert r.headers["x-export-matching-rows"] == str(N_ROWS)
    assert r.headers["x-export-rows"] == str(N_ROWS)
    assert r.headers["x-export-complete"] == "true"
    assert len({key(x) for x in body}) == N_ROWS
    src = {(a, b, c): v for a, b, c, v in _source_rows()}
    for x in body[::37]:
        k = (x[0].lstrip("'") if x[0].startswith("'") else x[0], x[1], x[2])
        v = src[k]
        assert (x[3] == "") if v is None else (float(x[3]) == pytest.approx(v))


def test_csv_filters_multiple_periods(client):
    periods = _periods()[10:14]
    f = json.dumps({"CATEGORY": ["Total", "a,b"], "TIME_PERIOD": periods})
    r = dl(client, filters=f)
    body = parse_csv(r)[1:]
    expect = sel(_source_rows(), {"Total", "a,b"}, set(periods))
    assert len(body) == len(expect) == 2 * 5 * 4
    assert sorted((x[0], x[1], x[2]) for x in body) == sorted(key(e) for e in expect)
    assert r.headers["x-export-matching-rows"] == str(len(expect))


def test_export_order_is_deterministic(client):
    a, b = dl(client).content, dl(client).content
    assert a == b


def test_xlsx_complete_with_filters(client, monkeypatch):
    f = json.dumps({"CATEGORY": ["Total", "a,b"], "REF_AREA": ["Cluj"]})
    r = dl(client, format="xlsx", filters=f)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content), read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("CATEGORY", "REF_AREA", "TIME_PERIOD", "OBS_VALUE")
    expect = [e for e in _source_rows() if e[0] in ("Total", "a,b") and e[1] == "Cluj"]
    assert len(rows) - 1 == len(expect) == 2 * N_PERIODS
    assert r.headers["x-export-rows"] == str(len(expect))
    got = {(x[0], x[1], x[2]): x[3] for x in rows[1:]}
    for c, a, t, v in expect[::41]:
        assert got[(c, a, t)] == ("" if v is None else pytest.approx(v))


def test_xlsx_full_selection_over_50k(client, monkeypatch):
    # whole 64k-row dataset fits a sheet: all rows must be present
    r = dl(client, format="xlsx")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content), read_only=True).active
    assert sum(1 for _ in ws.iter_rows(values_only=True)) == N_ROWS + 1


# ------------------------------------------------- 2. escaping / NULL / lang

def test_csv_roundtrip_labels_nulls_numbers(client):
    r = dl(client, safe=0)
    body = parse_csv(r)[1:]
    cats = {x[0] for x in body}
    assert cats == set(CATS)  # commas, newlines, quotes, unicode, formula-like verbatim
    assert any(x[3] == "" for x in body)  # NULL -> empty cell
    assert all(x[3] == "" or float(x[3]) is not None for x in body)


def test_csv_formula_safety_default_and_numbers_untouched(client):
    body = parse_csv(dl(client))[1:]
    cats = {x[0] for x in body}
    assert {"'=SUM(1)", "'+cmd", "'@x"} <= cats
    assert "-5" in cats                      # plain number is not a formula
    assert "=SUM(1)" not in cats
    assert all(not x[3].startswith("'") for x in body)  # OBS_VALUE never altered


def test_xlsx_formulas_stored_as_text_verbatim(client):
    f = json.dumps({"CATEGORY": ["=SUM(1)", "+cmd"], "REF_AREA": ["Cluj"],
                    "TIME_PERIOD": [_periods()[0]]})
    r = dl(client, format="xlsx", filters=f)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    vals = {ws.cell(i, 1).value: ws.cell(i, 1).data_type for i in (2, 3)}
    assert vals == {"+cmd": "s", "=SUM(1)": "s"}


def test_language_changes_labels_not_identity(client):
    f = json.dumps({"TIME_PERIOD": _periods()[:3]})
    ro = parse_csv(dl(client, filters=f))[1:]
    en = parse_csv(dl(client, filters=f, lang="en"))[1:]
    assert len(ro) == len(en) == len(CATS) * 5 * 3
    assert {x[0] for x in en} >= {"Total EN", "Country EN"}
    assert "Total" not in {x[0] for x in en}
    assert [x[1:] for x in ro] == [x[1:] for x in en] or \
        sorted(x[1:] for x in ro) == sorted(x[1:] for x in en)


# ------------------------------------------------------------- 3. XLSX policy

def test_xlsx_boundary_and_overflow(client, monkeypatch):
    import app.config as cfg
    f = json.dumps({"CATEGORY": ["Total"], "REF_AREA": ["Cluj"]})  # 1600 rows
    monkeypatch.setattr(cfg, "EXPORT_XLSX_MAX_ROWS", N_PERIODS)
    ok = dl(client, format="xlsx", filters=f)
    assert ok.status_code == 200
    ws = load_workbook(io.BytesIO(ok.content), read_only=True).active
    assert sum(1 for _ in ws.iter_rows(values_only=True)) == N_PERIODS + 1
    monkeypatch.setattr(cfg, "EXPORT_XLSX_MAX_ROWS", N_PERIODS - 1)
    before = set(glob.glob(os.path.join(tempfile.gettempdir(), "tempo-export-*")))
    bad = dl(client, format="xlsx", filters=f)
    assert bad.status_code == 413
    assert "CSV" in bad.json()["detail"]
    assert "content-disposition" not in bad.headers
    assert set(glob.glob(os.path.join(tempfile.gettempdir(), "tempo-export-*"))) == before
    # CSV for the same selection is unaffected
    assert len(parse_csv(dl(client, filters=f))) - 1 == N_PERIODS


def test_excel_constant():
    from app.services import export
    assert export.EXCEL_MAX_SHEET_ROWS == 1_048_576
    assert export.xlsx_row_limit() <= 1_048_575


def test_export_cap_config_rejects(client, monkeypatch):
    import app.config as cfg
    monkeypatch.setattr(cfg, "EXPORT_MAX_ROWS", 100)
    assert dl(client).status_code == 413
    assert dl(client, format="xlsx").status_code == 413


def test_preflight(client):
    r = dl(client, preflight=1, filters=json.dumps({"REF_AREA": ["Cluj"]}))
    assert r.status_code == 200
    assert r.json()["matching_rows"] == len(CATS) * N_PERIODS
    assert "content-disposition" not in r.headers


# --------------------------------------------------------------- 4xx handling

def test_errors(client):
    assert dl(client, code="NOPE").status_code == 404
    assert dl(client, filters="{bad").status_code == 400
    assert dl(client, filters=json.dumps({"NOCOL": ["x"]})).status_code == 400
    assert dl(client, format="pdf").status_code == 422
    assert dl(client, lang="fr").status_code == 422


# ------------------------------------------------ 4. memory and cleanup

def test_bounded_memory_csv_streaming(client):
    tracemalloc.start()
    with client.stream("GET", "/api/datasets/BIG1/download") as r:
        n = 0
        for chunk in r.iter_bytes():
            n += chunk.count(b"\n")
    cur, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    assert n > N_ROWS
    # whole CSV is ~3 MB; python-heap peak while streaming stays well below it
    # plus the (fixed) fixture/response overhead of the test client
    assert peak < 40 * 1024 * 1024


def test_xlsx_tempfile_removed_after_success_and_error(client, monkeypatch):
    pat = os.path.join(tempfile.gettempdir(), "tempo-export-*")
    before = set(glob.glob(pat))
    assert dl(client, format="xlsx", filters=json.dumps({"REF_AREA": ["Cluj"]})).status_code == 200
    assert set(glob.glob(pat)) == before
    from app.services import export
    monkeypatch.setattr(export, "make_row_transform",
                        lambda *a, **k: (lambda row: 1 / 0))
    assert dl(client, format="xlsx").status_code == 500
    assert set(glob.glob(pat)) == before


def test_generators_cleanup_on_early_close(tmp_path):
    from app.services import export
    con = duckdb.connect()
    cur = con.execute("SELECT range AS a FROM range(100000)")
    closed = []
    g = export.iter_csv(cur, ["a"], lambda r: list(r), 100, lambda: closed.append(1))
    next(g)
    g.close()                      # what Starlette does on client disconnect
    assert closed == [1]
    p = tmp_path / "f.bin"
    p.write_bytes(b"x" * 3_000_000)
    closed.clear()
    g = export.iter_file_then_delete(str(p), lambda: closed.append(1), chunk=1000)
    next(g)
    g.close()
    assert not p.exists() and closed == [1]
